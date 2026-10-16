"""
Match your own inventory (POS export) to scraped catalog products (name, image, URL).

Usage:  python match_products.py grocery_file.xlsx product_matches.xlsx
Needs:  pip install pandas openpyxl rapidfuzz

Edit ABBREV / WORDS below as you discover more POS shorthand, then re-run.
"""
import re, sys
import numpy as np
import pandas as pd
from rapidfuzz import fuzz, process

SRC = sys.argv[1] if len(sys.argv) > 1 else "grocery_file.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "product_matches.xlsx"

AUTO_SCORE = 88     # >= this with brand + size agreement -> auto-accept
REVIEW_SCORE = 70   # >= this -> goes to review sheet with top 3 candidates

# POS brand prefixes -> full brand (first word of your descriptions)
ABBREV = {
    "nat": "national", "p-f": "peek freans", "pf": "peek freans",
    "mit": "mitchells", "mitch": "mitchells", "a-g": "american garden",
    "h-s": "head shoulders", "m-r": "miss rose", "k&ns": "kns",
    "m-b": "murree brewery", "s-g": "saeed ghani", "bp": "bake parlor",
    "olper": "olpers", "salman": "salmans", "appolo": "apollo",
    "appollo": "apollo", "be": "be clean",
}
# Shorthand anywhere in the description
WORDS = {
    "b-s": "body spray", "f-w": "face wash", "powd": "powder", "pdr": "powder",
    "liq": "liquid", "shamp": "shampoo", "lrg": "large", "sml": "small",
    "med": "medium", "pch": "pouch", "kitchup": "ketchup", "tomoto": "tomato",
    "choc": "chocolate", "strawbery": "strawberry", "lavendor": "lavender",
    "opnr": "opener", "removel": "removal", "spary": "spray", "pcs": "pc",
}

UNIT = [  # (regex, canonical unit, multiplier to base)
    (r"kgs?|kilo", "g", 1000), (r"gms?|grams?|gr|g", "g", 1),
    (r"ltrs?|litres?|liters?|lt|l", "ml", 1000), (r"ml", "ml", 1),
]
SIZE_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(kgs?|kilo|gms?|grams?|gr|g|ltrs?|litres?|liters?|lt|l|ml)\b")


def sizes(s):
    out = set()
    for num, unit in SIZE_RE.findall(s):
        for pat, base, mult in UNIT:
            if re.fullmatch(pat, unit):
                out.add(f"{round(float(num) * mult, 1)}{base}")
                break
    return out


def norm(s, pos=False):
    s = str(s).lower().replace("'", "")
    if pos:
        first, *rest = s.split(" ", 1)
        if first in ABBREV:
            s = " ".join([ABBREV[first]] + rest)
        s = " ".join(WORDS.get(w, w) for w in s.split())
    s = re.sub(r"[^a-z0-9.& ]", " ", s)
    s = re.sub(r"(\d)\s+(kg|g|gm|ml|l|ltr)\b", r"\1\2", s)
    return " ".join(s.split())


VARIANT = {"zero", "diet", "lite", "light", "free", "men", "kids", "baby"}


def variant_ok(a, b):
    return (set(a.split()) & VARIANT) == (set(b.split()) & VARIANT)


def brand_ok(pos_norm, cand_norm):
    b = pos_norm.split()[0]
    head = " ".join(cand_norm.split()[:4])
    return len(b) > 1 and fuzz.partial_ratio(b, head) >= 85


def main():
    cat = pd.read_excel(SRC, sheet_name="Final")[["product_name", "image_url", "product_url"]]
    cat = cat.dropna(subset=["product_name"]).drop_duplicates("product_name").reset_index(drop=True)
    inv = pd.read_excel(SRC, sheet_name="kayanis data", header=1, dtype={"Item Code": str})
    inv = inv.dropna(subset=["Description"]).reset_index(drop=True)

    cat["n"] = cat.product_name.map(norm)
    cat["sz"] = cat.n.map(sizes)
    inv["n"] = inv.Description.map(lambda x: norm(x, pos=True))
    inv["sz"] = inv.n.map(sizes)

    scores = process.cdist(inv.n.tolist(), cat.n.tolist(), scorer=fuzz.token_set_ratio, workers=-1)
    scores2 = process.cdist(inv.n.tolist(), cat.n.tolist(), scorer=fuzz.token_sort_ratio, workers=-1)
    combo = 0.6 * scores + 0.4 * scores2

    rows = []
    for i, r in inv.iterrows():
        top = np.argsort(-combo[i])[:25]
        cands = []
        for j in top:
            c = cat.iloc[j]
            if not brand_ok(r.n, c.n) or not variant_ok(r.n, c.n):
                continue
            size_conflict = bool(r.sz and c.sz and not (r.sz & c.sz))
            if size_conflict:
                continue
            size_match = bool(r.sz and c.sz and (r.sz & c.sz))
            no_sizes = not r.sz and not c.sz
            cands.append((round(float(combo[i, j]), 1), size_match or no_sizes, j))
            if len(cands) == 3:
                break
        best = cands[0] if cands else None
        clear_winner = len(cands) < 2 or best[0] - cands[1][0] >= 3
        if best and best[0] >= AUTO_SCORE and best[1] and clear_winner:
            tier = "Auto"
        elif best and best[0] >= REVIEW_SCORE:
            tier = "Review"
        else:
            tier = "Unmatched"
        rec = {"Item Code": r["Item Code"], "Your Description": r.Description,
               "Units Sold (Aug)": r.QTY, "Price": r.PRICE, "Tier": tier}
        for k in range(3):
            if k < len(cands):
                s, sm, j = cands[k]
                c = cat.iloc[j]
                rec.update({f"C{k+1} Name": c.product_name, f"C{k+1} Score": s,
                            f"C{k+1} Image": c.image_url, f"C{k+1} URL": c.product_url})
        rows.append(rec)

    res = pd.DataFrame(rows).sort_values("Units Sold (Aug)", ascending=False)
    cat_brands = {n.split()[0] for n in cat.n if n}
    res["Why Unmatched"] = ""
    um = res.Tier == "Unmatched"
    res.loc[um, "Why Unmatched"] = [
        "Brand in catalog, no close product" if norm(d, pos=True).split()[0] in cat_brands
        else "Brand not in catalog" for d in res.loc[um, "Your Description"]]
    write_excel(res)
    return res


def write_excel(res):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    F = Font(name="Arial", size=10)
    H = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    HF = PatternFill("solid", fgColor="2F5233")
    INPUT = PatternFill("solid", fgColor="FFFF00")
    LINK = Font(name="Arial", size=10, color="0563C1", underline="single")

    wb = Workbook()
    summ = wb.active
    summ.title = "Summary"

    def sheet(name, df, cols, pick=False):
        ws = wb.create_sheet(name)
        headers = (["Pick (1/2/3/None)"] if pick else []) + cols
        ws.append(headers)
        for row in df[cols].itertuples(index=False):
            vals = [None if (isinstance(v, float) and pd.isna(v)) else v for v in row]
            ws.append(([None] if pick else []) + vals)
        for c in ws[1]:
            c.font, c.fill = H, HF
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.font = F
                h = headers[c.column - 1]
                if ("Image" in h or "URL" in h) and isinstance(c.value, str) and c.value.startswith("http"):
                    c.hyperlink, c.value, c.font = c.value, "open", LINK
        widths = {"Your Description": 38, "Pick (1/2/3/None)": 12}
        for i, h in enumerate(headers, 1):
            w = widths.get(h, 38 if "Name" in h else 16 if "Code" in h or "Why" in h else 9)
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "B2" if not pick else "C2"
        ws.auto_filter.ref = ws.dimensions
        if pick and ws.max_row > 1:
            dv = DataValidation(type="list", formula1='"1,2,3,None"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"A2:A{ws.max_row}")
            for r in range(2, ws.max_row + 1):
                ws.cell(r, 1).fill = INPUT
        return ws

    base = ["Item Code", "Your Description", "Units Sold (Aug)", "Price"]
    cands = [f"C{k} {f}" for k in (1, 2, 3) for f in ("Name", "Score", "Image", "URL")]
    sheet("Review", res[res.Tier == "Review"], base + cands, pick=True)
    sheet("Auto-Matched", res[res.Tier == "Auto"], base + ["C1 Name", "C1 Score", "C1 Image", "C1 URL"])
    sheet("Unmatched", res[res.Tier == "Unmatched"], base + ["Why Unmatched"])

    ab = wb.create_sheet("Abbreviations")
    ab.append(["POS shorthand", "Expands to", "Where"])
    for k, v in ABBREV.items():
        ab.append([k, v, "brand prefix"])
    for k, v in WORDS.items():
        ab.append([k, v, "anywhere"])
    for c in ab[1]:
        c.font, c.fill = H, HF
    for r in ab.iter_rows(min_row=2):
        for c in r:
            c.font = F
    ab.column_dimensions["A"].width = ab.column_dimensions["B"].width = 20

    summ.column_dimensions["A"].width = 44
    summ.column_dimensions["B"].width = 14
    summ.column_dimensions["C"].width = 18
    lines = [
        ("Product match summary", None, None),
        (None, None, None),
        ("Tier", "Items", "Units sold (Aug)"),
        ("Auto-matched (high confidence)", "=COUNTA('Auto-Matched'!B:B)-1", "=SUM('Auto-Matched'!C:C)"),
        ("Needs review (pick 1/2/3/None)", "=COUNTA(Review!C:C)-1", "=SUM(Review!D:D)"),
        ("Unmatched", "=COUNTA(Unmatched!B:B)-1", "=SUM(Unmatched!C:C)"),
        ("Total", "=SUM(B4:B6)", "=SUM(C4:C6)"),
        (None, None, None),
        ("Unmatched: brand not in catalog", '=COUNTIF(Unmatched!E:E,"Brand not in catalog")', None),
        ("Unmatched: brand in catalog, no close product", '=COUNTIF(Unmatched!E:E,"Brand in catalog, no close product")', None),
        (None, None, None),
        ("Review progress (picks entered)", "=COUNTA(Review!A:A)-1", None),
        (None, None, None),
        ("How to use: in Review, pick the correct candidate (yellow column). Sheets are sorted by units sold, so start at the top.", None, None),
        (f"Rules: Auto = score >= {AUTO_SCORE}, brand + size + variant agree, clear winner. Review = score >= {REVIEW_SCORE}.", None, None),
        ("Image links are scraped from punjabcashandcarry.com.pk and may need re-hosting before use.", None, None),
    ]
    for row in lines:
        summ.append(list(row))
    for r in summ.iter_rows():
        for c in r:
            c.font = F
    summ["A1"].font = Font(name="Arial", size=14, bold=True)
    for c in summ[3]:
        c.font, c.fill = H, HF
    for c in summ[7]:
        c.font = Font(name="Arial", size=10, bold=True)
    for r in range(4, 8):
        summ.cell(r, 3).number_format = "#,##0"
    wb.save(OUT)


if __name__ == "__main__":
    r = main()
    print(r.Tier.value_counts())
